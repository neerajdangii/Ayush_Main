from __future__ import annotations

from io import BytesIO

from django.core.files.base import ContentFile
from django.core.files.storage import default_storage
from django.utils import timezone

from .models import ReportRelease


RELEASE_COLUMNS = (
    ("S. No.", "serial_no"),
    ("Release Date", "release_date"),
    ("Analysis End Date", "analysis_end_date"),
    ("Previous Analysis End Date", "previous_analysis_end_date"),
    ("Released On (Date & Time)", "released_on"),
    ("Certificate No", "certificate_no"),
    ("Batch No", "batch_no"),
    ("Booking ID", "booking_code"),
    ("Sample Name", "sample_name"),
    ("Party", "party_name"),
    ("Channel", "channel"),
    ("Released By", "released_by_name"),
)


def _excel_text(value):
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _release_row(release, *, serial_no, analysis_end_date, previous_analysis_end_date, released_at):
    local_released_at = timezone.localtime(released_at) if released_at else None
    local_analysis_end_date = timezone.localtime(analysis_end_date) if analysis_end_date else None
    local_previous_analysis_end_date = (
        timezone.localtime(previous_analysis_end_date) if previous_analysis_end_date else None
    )
    row = []
    for _title, field in RELEASE_COLUMNS:
        if field == "serial_no":
            row.append(serial_no)
        elif field == "release_date":
            row.append(local_released_at.date() if local_released_at else "")
        elif field == "analysis_end_date":
            row.append(local_analysis_end_date.date() if local_analysis_end_date else "")
        elif field == "previous_analysis_end_date":
            row.append(
                local_previous_analysis_end_date.date() if local_previous_analysis_end_date else ""
            )
        elif field == "released_on":
            row.append(local_released_at.replace(tzinfo=None) if local_released_at else "")
        elif field == "channel":
            row.append(release.get_channel_display())
        else:
            row.append(_excel_text(getattr(release, field)))
    return row


def build_release_workbook(releases):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report Releases"
    sheet.freeze_panes = "A2"
    sheet.append([title for title, _field in RELEASE_COLUMNS])

    header_fill = PatternFill("solid", fgColor="2A6F8E")
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24

    def format_date_cells(row_index):
        sheet.cell(row_index, 2).number_format = "dd/mm/yyyy"
        sheet.cell(row_index, 3).number_format = "dd/mm/yyyy"
        sheet.cell(row_index, 4).number_format = "dd/mm/yyyy"
        sheet.cell(row_index, 5).number_format = "dd/mm/yyyy hh:mm:ss AM/PM"

    serial_no = 1
    for release in releases:
        if release.has_changes and release.previous_released_at:
            sheet.append(_release_row(
                release,
                serial_no=serial_no,
                analysis_end_date=release.previous_release_date,
                previous_analysis_end_date=None,
                released_at=release.previous_released_at,
            ))
            format_date_cells(sheet.max_row)
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor="FECACA")
            serial_no += 1

        sheet.append(_release_row(
            release,
            serial_no=serial_no,
            analysis_end_date=release.release_date,
            previous_analysis_end_date=(
                release.previous_release_date if release.analysis_end_date_changed else None
            ),
            released_at=release.released_at,
        ))
        format_date_cells(sheet.max_row)
        if release.analysis_end_date_changed:
            for cell in sheet[sheet.max_row]:
                cell.fill = PatternFill("solid", fgColor="FFF1F2")
        serial_no += 1

    widths = (10, 16, 18, 23, 25, 22, 20, 20, 34, 34, 16, 26)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def save_daily_release_workbook(day):
    releases = ReportRelease.objects.filter(released_at__date=day).order_by("released_at", "pk")
    content = build_release_workbook(releases)
    path = f"report_releases/daily/report-releases-{day.isoformat()}.xlsx"
    if default_storage.exists(path):
        default_storage.delete(path)
    return default_storage.save(path, ContentFile(content))
